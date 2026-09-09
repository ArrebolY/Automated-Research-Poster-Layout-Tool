from __future__ import annotations

import argparse
import base64
import csv
import html
import json
import os
import re
import shutil
import subprocess
import time
import urllib.error
import urllib.request
from dataclasses import dataclass
from pathlib import Path
from statistics import mean
from typing import Any

import fitz
from pptx import Presentation

from .batch_eval import _batch_stem, _slug, read_paper_list
from .config import LLMConfig
from .llm import TextGenerator, build_text_generator

EVALUATION_SOURCES = {
    "paperquiz": (
        "Paper2Poster: PaperQuiz evaluates whether a VLM reader can answer detail "
        "and understanding multiple-choice questions using only the rendered poster image."
    ),
    "faithfulness": (
        "FActScore-style factuality evaluation: decompose poster text into atomic factual claims "
        "and verify whether each claim is supported by the source paper. This measures hallucination "
        "or unsupported content rather than rewarding long posters."
    ),
    "visual_readability": (
        "Interim report: rule checks cover overflow, blank space, boundary issues, visual-text ratio, "
        "and reading order. Paper2Poster supports VLM-as-Judge over rendered poster images; this "
        "implementation can run the same style of visual judging through OpenAI vision models."
    ),
}

VLM_VISUAL_CRITERIA = [
    "aesthetic_element",
    "aesthetic_engagement",
    "aesthetic_layout",
    "information_low_level",
    "information_logic",
    "information_content",
]

@dataclass(slots=True)
class MethodSpec:
    name: str
    root: Path


@dataclass(slots=True)
class Question:
    id: str
    question: str
    options: list[str]
    answer: str
    aspect: str = ""


@dataclass(slots=True)
class VLMConfig:
    provider: str = "none"
    model: str = "gpt-4o-mini"
    api_key_env: str = "OPENAI_API_KEY"


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Evaluate final generated posters using the interim-report protocol: "
            "PaperQuiz-style QA, FActScore-style faithfulness, and visual readability checks."
        )
    )
    parser.add_argument("--paper-list", type=Path, required=True, help="Fixed list of paper.pdf paths.")
    parser.add_argument(
        "--method",
        action="append",
        default=[],
        metavar="NAME=ROOT",
        help=(
            "Poster output root for one method. If omitted, defaults to proposed, template_based, "
            "and direct_llm under outputs/."
        ),
    )
    parser.add_argument("--output-json", type=Path, default=Path("outputs/final_evaluation_summary.json"))
    parser.add_argument("--output-csv", type=Path, default=Path("outputs/final_evaluation_summary.csv"))
    parser.add_argument("--reports-root", type=Path, default=Path("outputs/assets"))
    parser.add_argument("--llm-provider", default="deepseek", choices=("rule", "openai", "deepseek", "openrouter", "gemini"))
    parser.add_argument("--llm-model", default="deepseek-v4-flash")
    parser.add_argument("--llm-api-key-env", default="")
    parser.add_argument("--llm-base-url", default="")
    parser.add_argument("--max-questions-per-type", type=int, default=20, help="0 means use all PaperQuiz questions.")
    parser.add_argument("--skip-paperquiz", action="store_true")
    parser.add_argument(
        "--paperquiz-evaluator",
        default="vlm",
        choices=("vlm", "text"),
        help="Use original Paper2Poster-style VLM PaperQuiz or the cheaper text-only proxy.",
    )
    parser.add_argument("--skip-faithfulness", action="store_true")
    parser.add_argument("--skip-fidelity", action="store_true", help="Deprecated alias for --skip-faithfulness.")
    parser.add_argument("--max-claims", type=int, default=12, help="Maximum poster factual claims to verify.")
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument(
        "--visual-evaluator",
        default="rule",
        choices=("rule", "vlm", "both"),
        help="Use rule checks, OpenAI VLM-as-Judge, or both for visual readability.",
    )
    parser.add_argument("--vlm-provider", default="none", choices=("none", "openai"))
    parser.add_argument("--vlm-model", default="gpt-4o-mini")
    parser.add_argument("--vlm-api-key-env", default="OPENAI_API_KEY")
    parser.add_argument("--poster-images-root", type=Path, default=Path("outputs/eval_images"))
    return parser


def main() -> None:
    args = build_parser().parse_args()
    paper_paths = read_paper_list(args.paper_list)
    if args.limit:
        paper_paths = paper_paths[: args.limit]

    methods = parse_methods(args.method)
    generator = build_text_generator(
        LLMConfig(
            provider=args.llm_provider,
            model=args.llm_model,
            api_key_env=args.llm_api_key_env,
            base_url=args.llm_base_url,
        )
    )
    skip_faithfulness = args.skip_faithfulness or args.skip_fidelity
    if generator is None and ((args.paperquiz_evaluator == "text" and not args.skip_paperquiz) or not skip_faithfulness):
        raise SystemExit(
            "Final text PaperQuiz/faithfulness evaluation needs an LLM. "
            "Use --skip-paperquiz --skip-faithfulness for visual-only checks."
        )
    vlm_config = VLMConfig(args.vlm_provider, args.vlm_model, args.vlm_api_key_env)
    if args.visual_evaluator in {"vlm", "both"} and vlm_config.provider == "none":
        raise SystemExit("--visual-evaluator vlm/both requires --vlm-provider openai.")
    if not args.skip_paperquiz and args.paperquiz_evaluator == "vlm" and vlm_config.provider != "openai":
        raise SystemExit("--paperquiz-evaluator vlm requires --vlm-provider openai.")

    results = []
    for paper_path in paper_paths:
        for method in methods:
            artifact = find_method_artifact(method, paper_path, args.paper_list.parent)
            report_dir = find_report_dir(args.reports_root, method.name, paper_path, args.paper_list.parent)
            result = evaluate_final_poster(
                paper_path=paper_path,
                method=method.name,
                artifact=artifact,
                report_dir=report_dir,
                generator=generator,
                max_questions_per_type=args.max_questions_per_type,
                skip_paperquiz=args.skip_paperquiz,
                paperquiz_evaluator=args.paperquiz_evaluator,
                skip_faithfulness=skip_faithfulness,
                max_claims=args.max_claims,
                visual_evaluator=args.visual_evaluator,
                vlm_config=vlm_config,
                poster_images_root=args.poster_images_root,
            )
            results.append(result)

    write_final_summary(results, args.output_json, args.output_csv)
    print(f"Evaluated {len(results)} poster artifacts.")
    print(f"Wrote {args.output_json}")
    print(f"Wrote {args.output_csv}")


def evaluate_final_poster(
    paper_path: Path,
    method: str,
    artifact: Path | None,
    report_dir: Path,
    generator: TextGenerator | None,
    max_questions_per_type: int = 10,
    skip_paperquiz: bool = False,
    paperquiz_evaluator: str = "vlm",
    skip_faithfulness: bool = False,
    max_claims: int = 12,
    visual_evaluator: str = "rule",
    vlm_config: VLMConfig | None = None,
    poster_images_root: Path = Path("outputs/eval_images"),
) -> dict:
    paper_id = _batch_stem(paper_path, paper_path.parents[1] if paper_path.name == "paper.pdf" else paper_path.parent)
    poster_text = extract_poster_text(artifact) if artifact and artifact.is_file() else ""
    paper_text = extract_paper_text(paper_path, report_dir)
    if not artifact or not artifact.is_file():
        paperquiz = {"status": "missing_artifact", "reason": "No poster artifact could be found."}
    elif skip_paperquiz:
        paperquiz = {"status": "skipped"}
    elif paperquiz_evaluator == "vlm":
        paperquiz = evaluate_paperquiz_vlm(
            paper_path.parent / "o3_qa.json",
            artifact,
            method=method,
            paper_id=paper_id,
            max_questions_per_type=max_questions_per_type,
            vlm_config=vlm_config or VLMConfig(),
            poster_images_root=poster_images_root,
        )
    elif not poster_text:
        paperquiz = {"status": "missing_artifact", "reason": "No poster text could be extracted."}
    else:
        paperquiz = evaluate_paperquiz(paper_path.parent / "o3_qa.json", poster_text, generator, max_questions_per_type)

    if not poster_text:
        faithfulness = {"status": "missing_artifact", "reason": "No poster text could be extracted."}
    else:
        faithfulness = (
            {"status": "skipped"}
            if skip_faithfulness
            else evaluate_claim_faithfulness(paper_text, poster_text, generator, max_claims=max_claims)
        )
    visual = evaluate_visual_readability(
        artifact,
        report_dir,
        poster_text,
        method=method,
        paper_id=paper_id,
        visual_evaluator=visual_evaluator,
        vlm_config=vlm_config,
        poster_images_root=poster_images_root,
    )
    overall = compute_overall_score(paperquiz, faithfulness, visual)
    result = {
        "paper_id": paper_id,
        "paper_path": str(paper_path),
        "method": method,
        "artifact": str(artifact) if artifact else "",
        "artifact_exists": bool(artifact and artifact.is_file()),
        "evaluation_sources": EVALUATION_SOURCES,
        "paperquiz": paperquiz,
        "faithfulness": faithfulness,
        "visual_readability": visual,
        "overall": overall,
    }
    output_dir = Path("outputs") / "final_eval" / method
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / f"{_slug(paper_id)}.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    return result


def parse_methods(raw_methods: list[str]) -> list[MethodSpec]:
    if not raw_methods:
        return [
            MethodSpec("proposed", Path("outputs/batch_posters/proposed")),
            MethodSpec("template_based", Path("outputs/batch_posters/template_based")),
            MethodSpec("direct_llm", Path("outputs/direct_llm")),
        ]
    methods = []
    for raw in raw_methods:
        if "=" not in raw:
            raise ValueError(f"Method must use NAME=ROOT format: {raw}")
        name, root = raw.split("=", 1)
        methods.append(MethodSpec(name.strip(), Path(root.strip())))
    return methods


def find_method_artifact(method: MethodSpec, paper_path: Path, root: Path) -> Path | None:
    base = _batch_stem(paper_path, root)
    mode_stem = base if method.name == "proposed" else f"{base}_{_slug(method.name, max_length=40)}"
    candidates = [
        method.root / f"{mode_stem}.pptx",
        method.root / f"{base}.pptx",
        method.root / mode_stem / "poster.html",
        method.root / base / "poster.html",
        method.root / f"{mode_stem}.html",
        method.root / f"{base}.html",
        Path("outputs") / "batch_posters" / method.name / f"{mode_stem}.pptx",
        Path("outputs") / "batch_posters" / method.name / f"{base}.pptx",
        Path("outputs") / "direct_llm" / mode_stem / "poster.html",
        Path("outputs") / "direct_llm" / base / "poster.html",
    ]
    for candidate in candidates:
        if candidate.is_file():
            return candidate
    return None


def find_report_dir(reports_root: Path, method: str, paper_path: Path, root: Path) -> Path:
    base = _batch_stem(paper_path, root)
    mode_stem = base if method == "proposed" else f"{base}_{_slug(method, max_length=40)}"
    for candidate in (reports_root / mode_stem, reports_root / base):
        if candidate.is_dir():
            return candidate
    return reports_root / mode_stem


def extract_poster_text(artifact: Path) -> str:
    suffix = artifact.suffix.lower()
    if suffix == ".pptx":
        return extract_pptx_text(artifact)
    if suffix in {".html", ".htm"}:
        return extract_html_text(artifact)
    if suffix == ".txt":
        return artifact.read_text(encoding="utf-8", errors="replace")
    return ""


def extract_pptx_text(path: Path) -> str:
    prs = Presentation(path)
    chunks = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                chunks.append(shape.text)
    return normalize_space("\n".join(chunks))


def extract_html_text(path: Path) -> str:
    raw = path.read_text(encoding="utf-8", errors="replace")
    raw = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", raw)
    raw = re.sub(r"(?is)<[^>]+>", " ", raw)
    return normalize_space(html.unescape(raw))


def extract_paper_text(paper_path: Path, report_dir: Path) -> str:
    markdown_candidates = [
        report_dir / "paper.md",
        report_dir / "paper_markdown.md",
        report_dir / f"{paper_path.stem}.md",
    ]
    for candidate in markdown_candidates:
        if candidate.is_file():
            return normalize_space(candidate.read_text(encoding="utf-8", errors="replace"))
    document = fitz.open(paper_path)
    return normalize_space("\n".join(page.get_text("text") for page in document))


def evaluate_paperquiz(
    qa_path: Path,
    poster_text: str,
    generator: TextGenerator | None,
    max_questions_per_type: int,
) -> dict:
    if generator is None:
        return {"status": "skipped", "reason": "no_llm"}
    if not qa_path.is_file():
        return {"status": "missing_qa", "qa_path": str(qa_path)}
    qa = json.loads(qa_path.read_text(encoding="utf-8"))
    detail = normalize_questions(qa.get("detail", {}), "detail", max_questions_per_type)
    understanding = normalize_questions(qa.get("understanding", {}), "understanding", max_questions_per_type)
    detail_result = answer_questions_with_llm(poster_text, detail, generator)
    understanding_result = answer_questions_with_llm(poster_text, understanding, generator)
    return {
        "status": "ok",
        "source": "Paper2Poster PaperQuiz / o3_qa.json",
        "qa_path": str(qa_path),
        "detail": detail_result,
        "understanding": understanding_result,
        "average_accuracy": round(mean([detail_result["accuracy"], understanding_result["accuracy"]]), 3),
    }


def evaluate_paperquiz_vlm(
    qa_path: Path,
    artifact: Path,
    method: str,
    paper_id: str,
    max_questions_per_type: int,
    vlm_config: VLMConfig,
    poster_images_root: Path,
) -> dict:
    if vlm_config.provider != "openai":
        return {"status": "unsupported_vlm_provider", "provider": vlm_config.provider}
    if not qa_path.is_file():
        return {"status": "missing_qa", "qa_path": str(qa_path)}
    image_path = poster_images_root / method / f"{_slug(paper_id)}.png"
    render_result = render_artifact_to_png(artifact, image_path)
    if render_result.get("status") != "ok":
        return {
            "status": "render_failed",
            "qa_path": str(qa_path),
            "artifact": str(artifact),
            "render": render_result,
        }
    qa = json.loads(qa_path.read_text(encoding="utf-8"))
    detail = normalize_questions(qa.get("detail", {}), "detail", max_questions_per_type)
    understanding = normalize_questions(qa.get("understanding", {}), "understanding", max_questions_per_type)
    try:
        detail_result = answer_questions_with_vlm(image_path, detail, vlm_config)
        understanding_result = answer_questions_with_vlm(image_path, understanding, vlm_config)
    except Exception as exc:
        return {
            "status": "vlm_failed",
            "source": "Original Paper2Poster-style PaperQuiz: VLM answers o3_qa.json questions from the rendered poster image.",
            "qa_path": str(qa_path),
            "artifact_image": str(image_path),
            "model": vlm_config.model,
            "error": f"{type(exc).__name__}: {exc}",
        }
    return {
        "status": "ok",
        "source": "Original Paper2Poster-style PaperQuiz: VLM answers o3_qa.json questions from the rendered poster image.",
        "qa_path": str(qa_path),
        "artifact_image": str(image_path),
        "model": vlm_config.model,
        "detail": detail_result,
        "understanding": understanding_result,
        "average_accuracy": round(mean([detail_result["accuracy"], understanding_result["accuracy"]]), 3),
    }


def normalize_questions(block: dict, label: str, max_questions: int) -> list[Question]:
    questions = block.get("questions", {}) or {}
    answers = block.get("answers", {}) or {}
    aspects = block.get("aspects", {}) or {}
    items = list(questions.items()) if isinstance(questions, dict) else list(enumerate(questions, start=1))
    if max_questions > 0:
        items = items[:max_questions]
    result = []
    for index, (raw_id, raw_question) in enumerate(items, start=1):
        question_id = str(raw_id)
        if isinstance(raw_question, dict):
            question_text = str(raw_question.get("question", ""))
            options = [str(option) for option in raw_question.get("options", [])]
        else:
            question_text = str(raw_question)
            options = []
        answer = answers.get(raw_id, answers.get(question_id, answers.get(f"Question {index}", "")))
        aspect = aspects.get(raw_id, aspects.get(question_id, aspects.get(f"Question {index}", "")))
        result.append(
            Question(
                id=f"{label}_{index}",
                question=question_text,
                options=options,
                answer=str(answer),
                aspect=str(aspect),
            )
        )
    return result


def answer_questions_with_vlm(image_path: Path, questions: list[Question], config: VLMConfig) -> dict:
    if not questions:
        return {"question_count": 0, "correct_count": 0, "accuracy": 0.0, "aspect_accuracy": {}, "answers": []}
    payload = [{"id": q.id, "question": q.question, "options": q.options} for q in questions]
    response_data = openai_vlm_answer_questions(image_path, payload, config)
    parsing_error = ""
    try:
        predicted = {
            str(item.get("id", "")): str(item.get("answer", ""))
            for item in response_data.get("answers", [])
        }
    except AttributeError as exc:
        parsing_error = f"{type(exc).__name__}: {exc}"
        predicted = {}
    if not predicted:
        predicted = extract_qa_answers_from_text(json.dumps(response_data, ensure_ascii=False))
    return score_predicted_answers(predicted, questions, parsing_error=parsing_error)


def answer_questions_with_llm(poster_text: str, questions: list[Question], generator: TextGenerator) -> dict:
    if not questions:
        return {"question_count": 0, "correct_count": 0, "accuracy": 0.0, "aspect_accuracy": {}, "answers": []}
    payload = [
        {"id": q.id, "question": q.question, "options": q.options}
        for q in questions
    ]
    prompt = f"""
You are evaluating a generated academic poster using the PaperQuiz idea.
Use ONLY the poster text below as evidence. Answer each multiple-choice question.
If the poster does not contain enough information, choose the closest answer but mark low confidence.

Return only valid JSON. Do not include evidence, explanations, markdown, or comments.
Use this exact format:
{{"answers":[{{"id":"detail_1","answer":"A"}}]}}

Poster text:
{poster_text[:12000]}

Questions:
{json.dumps(payload, ensure_ascii=False)}
""".strip()
    response = generator.generate(prompt)
    parsing_error = ""
    try:
        data = extract_json_object(response)
        predicted = {str(item.get("id", "")): str(item.get("answer", "")) for item in data.get("answers", [])}
    except (json.JSONDecodeError, ValueError) as exc:
        parsing_error = f"{type(exc).__name__}: {exc}"
        predicted = extract_qa_answers_from_text(response)
    return score_predicted_answers(predicted, questions, parsing_error=parsing_error)


def score_predicted_answers(predicted: dict[str, str], questions: list[Question], parsing_error: str = "") -> dict:
    answer_rows = []
    aspect_totals: dict[str, list[bool]] = {}
    for question in questions:
        predicted_letter = answer_letter(predicted.get(question.id, ""))
        reference_letter = answer_letter(question.answer)
        is_correct = bool(predicted_letter and reference_letter and predicted_letter == reference_letter)
        aspect = question.aspect or "unknown"
        aspect_totals.setdefault(aspect, []).append(is_correct)
        answer_rows.append(
            {
                "id": question.id,
                "question": question.question,
                "predicted": predicted_letter,
                "reference": reference_letter,
                "correct": is_correct,
                "aspect": aspect,
            }
        )
    correct = sum(1 for row in answer_rows if row["correct"])
    return {
        "question_count": len(questions),
        "correct_count": correct,
        "accuracy": round(correct / len(questions), 3),
        "parsing_error": parsing_error,
        "aspect_accuracy": {
            aspect: round(sum(values) / len(values), 3)
            for aspect, values in sorted(aspect_totals.items())
        },
        "answers": answer_rows,
    }


def evaluate_claim_faithfulness(
    paper_text: str,
    poster_text: str,
    generator: TextGenerator | None,
    max_claims: int = 12,
) -> dict:
    if generator is None:
        return {"status": "skipped", "reason": "no_llm"}
    prompt = f"""
You are evaluating factual faithfulness of a generated academic poster.
Use a FActScore-style procedure:
1. Extract up to {max_claims} atomic factual claims from the poster text.
2. For each claim, decide whether it is supported by the original paper text.

Labels:
- supported: the paper text clearly supports the claim.
- unsupported: the claim contradicts the paper or cannot be found in the paper.
- uncertain: the paper text is insufficient or ambiguous.

Important:
- Do not reward longer posters.
- Do not penalize missing information. Only judge claims that the poster actually makes.
- A concise poster can score highly if its claims are correct.
- Be strict with numbers, datasets, model names, author claims, and conclusions.

Return only valid JSON:
{{
  "claims": [
    {{"claim": "...", "label": "supported"}}
  ]
}}

Do not include evidence, explanations, markdown, or comments in the JSON.
Use only the labels: supported, unsupported, uncertain.

Original paper text:
{paper_text[:16000]}

Generated poster text:
{poster_text[:12000]}
""".strip()
    response = generator.generate(prompt)
    parsing_error = ""
    try:
        data = extract_json_object(response)
        raw_claims = data.get("claims", [])
    except (json.JSONDecodeError, ValueError) as exc:
        parsing_error = f"{type(exc).__name__}: {exc}"
        raw_claims = extract_claims_from_text(response)
    claims = normalize_claim_judgments(raw_claims)
    supported = sum(1 for item in claims if item["label"] == "supported")
    unsupported = sum(1 for item in claims if item["label"] == "unsupported")
    uncertain = sum(1 for item in claims if item["label"] == "uncertain")
    total = len(claims)
    score = round(supported / total * 100, 3) if total else 0.0
    return {
        "status": "ok" if total else "no_claims",
        "source": "FActScore-style atomic claim verification against the source paper.",
        "claim_count": total,
        "supported_claim_count": supported,
        "unsupported_claim_count": unsupported,
        "uncertain_claim_count": uncertain,
        "faithfulness_score": score,
        "unsupported_claim_rate": round(unsupported / total, 3) if total else 0.0,
        "uncertain_claim_rate": round(uncertain / total, 3) if total else 0.0,
        "parsing_error": parsing_error,
        "claims": claims,
    }


def evaluate_visual_readability(
    artifact: Path | None,
    report_dir: Path,
    poster_text: str,
    method: str = "unknown",
    paper_id: str = "paper",
    visual_evaluator: str = "rule",
    vlm_config: VLMConfig | None = None,
    poster_images_root: Path = Path("outputs/eval_images"),
) -> dict:
    rule = evaluate_rule_visual_readability(artifact, report_dir, poster_text)
    if visual_evaluator == "rule":
        return rule

    vlm = evaluate_vlm_visual_readability(
        artifact=artifact,
        method=method,
        paper_id=paper_id,
        vlm_config=vlm_config or VLMConfig(),
        poster_images_root=poster_images_root,
    )
    if visual_evaluator == "vlm":
        return vlm

    combined = dict(rule)
    combined["source"] = "Combined rule checks and OpenAI VLM-as-Judge over rendered poster image."
    combined["rule"] = rule
    combined["vlm"] = vlm
    if vlm.get("status") == "ok":
        combined["status"] = "ok"
        combined["rule_score"] = rule.get("score", 0.0)
        combined["vlm_score"] = vlm.get("score", 0.0)
        combined["score"] = combined_visual_score(rule, vlm)
        combined["score_source"] = "vlm_rule_combined"
        combined["score_weights"] = {"vlm": 0.7, "rule": 0.3}
    else:
        combined["rule_score"] = rule.get("score", 0.0)
        combined["vlm_score"] = 0.0
        combined["score_source"] = "rule_fallback"
    return combined


def combined_visual_score(rule: dict, vlm: dict) -> float:
    """Combine holistic VLM judgment with deterministic layout defect checks."""
    rule_score = float(rule.get("score", 0.0) or 0.0)
    vlm_score = float(vlm.get("score", 0.0) or 0.0)
    return round(clamp(vlm_score * 0.7 + rule_score * 0.3, 1, 5), 3)


def evaluate_rule_visual_readability(artifact: Path | None, report_dir: Path, poster_text: str) -> dict:
    if artifact and artifact.suffix.lower() == ".pptx":
        return evaluate_visual_from_reports(report_dir, poster_text)
    if artifact and artifact.suffix.lower() in {".html", ".htm"}:
        return evaluate_visual_from_html(artifact, poster_text)
    return {"status": "missing_artifact", "score": 0.0}


def evaluate_visual_from_reports(report_dir: Path, poster_text: str) -> dict:
    panel = read_json(report_dir / "panel_quality.json")
    layout = read_json(report_dir / "layout_quality.json")
    summary = panel.get("summary", {}) or {}
    issues = summary.get("issue_counts", {}) or {}
    overflow = int(issues.get("text_overflow_risk", 0) or 0)
    blank = int(issues.get("excessive_blank_space", 0) or 0) + int(issues.get("pure_text_too_spacious", 0) or 0)
    small_visual = int(issues.get("image_too_small", 0) or 0) + int(issues.get("table_too_small", 0) or 0)
    out_of_bounds = int(issues.get("out_of_bounds", 0) or 0)
    panel_count = max(1, len(panel.get("panels", []) or []))
    figure_count = len((layout.get("figures", []) or [])) if layout else 0
    base = 5.0
    score = base - overflow * 0.9 - out_of_bounds * 1.0 - blank * 0.35 - small_visual * 0.35
    visual_text_ratio = round(figure_count / panel_count, 3)
    return {
        "status": "ok",
        "source": "Interim-report rule checks; Paper2Poster VLM-as-Judge aspects are approximated without a vision model.",
        "score": round(clamp(score, 1, 5), 3),
        "text_overflow_count": overflow,
        "blank_space_count": blank,
        "small_visual_count": small_visual,
        "out_of_bounds_count": out_of_bounds,
        "figure_count": figure_count,
        "panel_count": panel_count,
        "poster_word_count": count_words(poster_text),
        "visual_text_ratio": visual_text_ratio,
        "subscores": {
            "aesthetic_layout": round(clamp(5 - blank * 0.5 - out_of_bounds, 1, 5), 3),
            "information_low_level": round(clamp(5 - overflow - small_visual * 0.35, 1, 5), 3),
            "information_content_proxy": round(clamp(2 + min(3, count_words(poster_text) / 180), 1, 5), 3),
            "visual_text_ratio_proxy": round(clamp(3 + min(2, visual_text_ratio), 1, 5), 3),
        },
    }


def evaluate_visual_from_html(artifact: Path, poster_text: str) -> dict:
    raw = artifact.read_text(encoding="utf-8", errors="replace")
    image_paths = re.findall(r"<img[^>]+src=[\"']([^\"']+)", raw, flags=re.I)
    missing = []
    for src in image_paths:
        if src.startswith("file://"):
            if not Path(src[7:]).is_file():
                missing.append(src)
        elif not src.startswith(("http://", "https://", "data:")):
            if not (artifact.parent / src).is_file():
                missing.append(src)
    word_count = count_words(poster_text)
    score = 5.0
    if not image_paths:
        score -= 1.5
    score -= min(2.0, len(missing) * 0.8)
    if word_count < 180:
        score -= 0.6
    return {
        "status": "ok",
        "source": "HTML readability proxy for direct LLM baseline.",
        "score": round(clamp(score, 1, 5), 3),
        "image_count": len(image_paths),
        "missing_image_count": len(missing),
        "poster_word_count": word_count,
        "subscores": {
            "aesthetic_layout": round(clamp(score, 1, 5), 3),
            "information_low_level": round(clamp(5 if 180 <= word_count <= 900 else 4, 1, 5), 3),
            "information_content_proxy": round(clamp(2 + min(3, word_count / 180), 1, 5), 3),
            "visual_text_ratio_proxy": round(clamp(2 + min(3, len(image_paths) * 0.7), 1, 5), 3),
        },
    }


def evaluate_vlm_visual_readability(
    artifact: Path | None,
    method: str,
    paper_id: str,
    vlm_config: VLMConfig,
    poster_images_root: Path,
) -> dict:
    if not artifact or not artifact.is_file():
        return {"status": "missing_artifact", "score": 0.0}
    if vlm_config.provider != "openai":
        return {"status": "unsupported_vlm_provider", "provider": vlm_config.provider, "score": 0.0}
    image_dir = poster_images_root / method
    image_path = image_dir / f"{_slug(paper_id)}.png"
    render_result = render_artifact_to_png(artifact, image_path)
    if render_result.get("status") != "ok":
        return {
            "status": "render_failed",
            "score": 0.0,
            "artifact": str(artifact),
            "render": render_result,
        }
    try:
        judge = openai_vlm_judge(image_path, vlm_config)
    except Exception as exc:
        return {
            "status": "vlm_failed",
            "score": 0.0,
            "artifact": str(artifact),
            "image_path": str(image_path),
            "error": f"{type(exc).__name__}: {exc}",
        }
    return {
        "status": "ok",
        "source": "OpenAI VLM-as-Judge over a rendered poster image, following Paper2Poster-style visual judging.",
        "score": judge["score"],
        "image_path": str(image_path),
        "model": vlm_config.model,
        "criteria": judge["criteria"],
        "aesthetic_average": judge.get("aesthetic_average", 0.0),
        "information_average": judge.get("information_average", 0.0),
        "overall": judge.get("overall", {}),
    }


def render_artifact_to_png(artifact: Path, output_path: Path) -> dict:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.is_file() and output_path.stat().st_size > 0:
        return {"status": "ok", "backend": "cached", "image_path": str(output_path)}
    suffix = artifact.suffix.lower()
    if suffix in {".png", ".jpg", ".jpeg"}:
        shutil.copy2(artifact, output_path)
        return {"status": "ok", "backend": "copy", "image_path": str(output_path)}
    return render_with_qlmanage(artifact, output_path)


def render_with_qlmanage(artifact: Path, output_path: Path) -> dict:
    qlmanage = shutil.which("qlmanage")
    if not qlmanage:
        return {"status": "missing_renderer", "reason": "qlmanage is not available on this machine."}
    temp_dir = output_path.parent / f".{output_path.stem}_ql"
    if temp_dir.exists():
        shutil.rmtree(temp_dir)
    temp_dir.mkdir(parents=True, exist_ok=True)
    command = [qlmanage, "-t", "-s", "1800", "-o", str(temp_dir), str(artifact)]
    completed = subprocess.run(command, capture_output=True, text=True, timeout=90, check=False)
    if completed.returncode != 0:
        shutil.rmtree(temp_dir, ignore_errors=True)
        return {
            "status": "renderer_error",
            "backend": "qlmanage",
            "returncode": completed.returncode,
            "stderr": completed.stderr[-1000:],
            "stdout": completed.stdout[-1000:],
        }
    candidates = sorted(temp_dir.glob("*.png"), key=lambda path: path.stat().st_size if path.exists() else 0, reverse=True)
    if not candidates:
        shutil.rmtree(temp_dir, ignore_errors=True)
        return {"status": "renderer_error", "backend": "qlmanage", "reason": "No PNG thumbnail was produced."}
    shutil.move(str(candidates[0]), output_path)
    shutil.rmtree(temp_dir, ignore_errors=True)
    return {"status": "ok", "backend": "qlmanage", "image_path": str(output_path)}


def openai_vlm_judge(image_path: Path, config: VLMConfig) -> dict:
    api_key = os.getenv(config.api_key_env, "").strip()
    if not api_key:
        raise RuntimeError(f"Missing API key. Set the {config.api_key_env} environment variable.")
    image_data = base64.b64encode(image_path.read_bytes()).decode("ascii")
    payload = {
        "model": config.model,
        "messages": [
            {
                "role": "system",
                "content": (
                    "You are an expert evaluator for academic research posters. "
                    "Judge the rendered poster image visually and return strict JSON only."
                ),
            },
            {
                "role": "user",
                "content": [
                    {
                        "type": "text",
                        "text": visual_judge_prompt(),
                    },
                    {
                        "type": "image_url",
                        "image_url": {"url": f"data:image/png;base64,{image_data}"},
                    },
                ],
            },
        ],
        "temperature": 0,
        "max_tokens": 900,
        "response_format": {"type": "json_object"},
    }
    data = post_openai_json("https://api.openai.com/v1/chat/completions", payload, api_key)
    content = data["choices"][0]["message"].get("content", "")
    parsed = extract_json_object(content)
    return normalize_vlm_judge_result(parsed)


def openai_vlm_answer_questions(image_path: Path, questions_payload: list[dict[str, Any]], config: VLMConfig) -> dict:
    api_key = os.getenv(config.api_key_env, "").strip()
    if not api_key:
        raise RuntimeError(f"Missing API key. Set the {config.api_key_env} environment variable.")
    image_data = base64.b64encode(image_path.read_bytes()).decode("ascii")
    prompt = f"""
You are answering PaperQuiz multiple-choice questions using ONLY the rendered academic poster image.
Read the poster visually, including headings, text, figures, tables, captions, and labels.
If the answer is not visible or not inferable from the poster image, choose the closest option, but do not use outside knowledge.

Return only valid JSON:
{{"answers":[{{"id":"detail_1","answer":"A"}}]}}

Questions:
{json.dumps(questions_payload, ensure_ascii=False)}
""".strip()
    payload = {
        "model": config.model,
        "messages": [
            {
                "role": "system",
                "content": (
                    "You are an evaluator for PaperQuiz. Answer multiple-choice questions "
                    "from the poster image only and return strict JSON."
                ),
            },
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": prompt},
                    {
                        "type": "image_url",
                        "image_url": {"url": f"data:image/png;base64,{image_data}"},
                    },
                ],
            },
        ],
        "temperature": 0,
        "max_tokens": 1600,
        "response_format": {"type": "json_object"},
    }
    data = post_openai_json("https://api.openai.com/v1/chat/completions", payload, api_key)
    content = data["choices"][0]["message"].get("content", "")
    return extract_json_object(content)


def visual_judge_prompt() -> str:
    return """
Evaluate this academic research poster image as a strict conference-poster reviewer.
Do not reward the poster simply because it is readable. A normal, acceptable poster should receive 3.
Use 4 only when the design is clearly strong and professional. Use 5 only for near-publication-quality posters.
If you see large unused regions, tiny visuals, crowded text, weak hierarchy, broken images, clipping, or awkward reading order, deduct points.

Score each criterion from 1 to 5:
1 = serious problem that would be unacceptable in a presentation poster
2 = weak; multiple visible issues
3 = acceptable but ordinary; some issues remain
4 = good; only minor issues
5 = excellent; polished and highly effective

Criteria:
- aesthetic_element: typography, color palette, borders, panels, and visual design elements.
- aesthetic_engagement: whether the poster looks professionally engaging rather than plain, messy, or unfinished.
- aesthetic_layout: balance, alignment, spacing, use of white space, and absence of crowding.
- information_low_level: text readability, image/table clarity, no clipping, no broken images, and no overlaps.
- information_logic: reading order, section hierarchy, title prominence, and coherent flow.
- information_content: whether visible text and visuals jointly explain the paper rather than looking generic or disconnected.

Calibration examples:
- A poster with a huge empty lower half should not score above 3 for aesthetic_layout.
- A poster with clear images but tiny text should not score above 3 for information_low_level.
- A poster with broken image icons should score 1 or 2 for information_low_level.
- A visually acceptable but plain template should usually be around 3, not 4 or 5.
- A poster whose figures are far from their related explanations should lose points on information_content and information_logic.

Return only JSON:
{
  "criteria": {
    "aesthetic_element": {"score": 1, "reason": "..."},
    "aesthetic_engagement": {"score": 1, "reason": "..."},
    "aesthetic_layout": {"score": 1, "reason": "..."},
    "information_low_level": {"score": 1, "reason": "..."},
    "information_logic": {"score": 1, "reason": "..."},
    "information_content": {"score": 1, "reason": "..."}
  },
  "overall": {"score": 1, "reason": "..."}
}
""".strip()


def post_openai_json(url: str, payload: dict, api_key: str) -> dict:
    body = json.dumps(payload).encode("utf-8")
    last_error: Exception | None = None
    for attempt in range(7):
        request = urllib.request.Request(
            url,
            data=body,
            headers={"Content-Type": "application/json", "Authorization": f"Bearer {api_key}"},
            method="POST",
        )
        try:
            with urllib.request.urlopen(request, timeout=90) as response:
                return json.loads(response.read().decode("utf-8"))
        except urllib.error.HTTPError as exc:
            details = exc.read().decode("utf-8", errors="replace")
            if exc.code not in {408, 409, 425, 429, 500, 502, 503, 504}:
                raise RuntimeError(f"OpenAI VLM request failed with HTTP {exc.code}: {details}") from exc
            last_error = RuntimeError(f"OpenAI VLM request failed with HTTP {exc.code}: {details}")
        except (TimeoutError, urllib.error.URLError, OSError) as exc:
            last_error = exc
        if attempt < 6:
            time.sleep(retry_sleep_seconds(last_error, attempt))
    raise RuntimeError(f"OpenAI VLM request failed after retries: {type(last_error).__name__}: {last_error}") from last_error


def retry_sleep_seconds(error: Exception | None, attempt: int) -> float:
    message = str(error or "")
    wait_match = re.search(r"try again in\s+(\d+)ms", message, flags=re.I)
    provider_wait = float(wait_match.group(1)) / 1000 if wait_match else 0.0
    if "429" in message or "rate_limit" in message.lower():
        return max(provider_wait + 1.0, 5.0 * (attempt + 1))
    return max(provider_wait + 0.5, 2.0 * (attempt + 1))


def normalize_vlm_judge_result(data: dict) -> dict:
    raw_criteria = data.get("criteria", {}) or {}
    criteria = {}
    scores = []
    for name in VLM_VISUAL_CRITERIA:
        item = raw_criteria.get(name, {}) if isinstance(raw_criteria, dict) else {}
        score = int(clamp(float(item.get("score", 0) or 0), 1, 5))
        scores.append(score)
        criteria[name] = {
            "score": score,
            "reason": str(item.get("reason", "")),
        }
    overall = data.get("overall", {}) or {}
    overall_score = overall.get("score", None)
    score = float(overall_score) if overall_score is not None else mean(scores)
    aesthetic = [criteria[name]["score"] for name in VLM_VISUAL_CRITERIA if name.startswith("aesthetic_")]
    information = [criteria[name]["score"] for name in VLM_VISUAL_CRITERIA if name.startswith("information_")]
    return {
        "criteria": criteria,
        "aesthetic_average": round(mean(aesthetic), 3) if aesthetic else 0.0,
        "information_average": round(mean(information), 3) if information else 0.0,
        "overall": {
            "score": round(clamp(score, 1, 5), 3),
            "reason": str(overall.get("reason", "")),
        },
        "score": round(mean(scores), 3),
    }


def compute_overall_score(paperquiz: dict, faithfulness: dict, visual: dict) -> dict:
    components = {}
    if paperquiz.get("status") == "ok":
        components["paperquiz_accuracy"] = float(paperquiz.get("average_accuracy", 0.0)) * 100
    if faithfulness.get("status") == "ok":
        components["faithfulness_score"] = float(faithfulness.get("faithfulness_score", 0.0))
    if visual.get("status") == "ok":
        components["visual_readability_score"] = float(visual.get("score", 0.0)) / 5 * 100
    return {
        "component_scores": components,
        "overall_score": round(mean(components.values()), 3) if components else 0.0,
    }


def write_final_summary(results: list[dict], output_json: Path, output_csv: Path) -> None:
    output_json.parent.mkdir(parents=True, exist_ok=True)
    output_csv.parent.mkdir(parents=True, exist_ok=True)
    output_json.write_text(
        json.dumps(
            {
                "poster_count": len(results),
                "evaluation_sources": EVALUATION_SOURCES,
                "aggregate": aggregate_final_results(results),
                "posters": results,
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    rows = [flatten_final_result(result) for result in results]
    fieldnames = sorted({key for row in rows for key in row})
    with output_csv.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def aggregate_final_results(results: list[dict]) -> dict:
    by_method: dict[str, list[dict]] = {}
    for result in results:
        by_method.setdefault(result.get("method", ""), []).append(flatten_final_result(result))
    aggregate = {}
    for method, rows in by_method.items():
        numeric_keys = sorted(
            key
            for key in {key for row in rows for key in row}
            if all(isinstance(row.get(key), int | float | bool) for row in rows if key in row)
        )
        aggregate[method] = {
            f"avg_{key}": round(mean(float(row[key]) for row in rows if key in row), 3)
            for key in numeric_keys
            if any(key in row for row in rows)
        }
    return aggregate


def flatten_final_result(result: dict) -> dict:
    row = {
        "paper_id": result.get("paper_id", ""),
        "method": result.get("method", ""),
        "artifact": result.get("artifact", ""),
        "artifact_exists": result.get("artifact_exists", False),
    }
    paperquiz = result.get("paperquiz", {}) or {}
    faithfulness = result.get("faithfulness", {}) or {}
    visual = result.get("visual_readability", {}) or {}
    vlm_visual = visual.get("vlm", visual) if isinstance(visual.get("vlm", visual), dict) else {}
    vlm_criteria = vlm_visual.get("criteria", {}) or {}
    overall = result.get("overall", {}) or {}
    row.update(
        {
            "paperquiz.status": paperquiz.get("status", ""),
            "paperquiz.ok": paperquiz.get("status") == "ok",
            "faithfulness.status": faithfulness.get("status", ""),
            "faithfulness.ok": faithfulness.get("status") == "ok",
            "visual_readability.status": visual.get("status", ""),
            "visual_readability.score_source": visual.get("score_source", ""),
            "visual_readability.vlm_status": vlm_visual.get("status", ""),
            "visual_readability.ok": visual.get("status") == "ok",
            "visual_readability.vlm_ok": vlm_visual.get("status") == "ok",
            "overall.overall_score": overall.get("overall_score", 0.0),
        }
    )
    if paperquiz.get("status") == "ok":
        row.update(
            {
                "paperquiz.average_accuracy": paperquiz.get("average_accuracy", 0.0),
                "paperquiz.detail_accuracy": (paperquiz.get("detail", {}) or {}).get("accuracy", 0.0),
                "paperquiz.understanding_accuracy": (paperquiz.get("understanding", {}) or {}).get("accuracy", 0.0),
            }
        )
    if faithfulness.get("status") == "ok":
        row.update(
            {
                "faithfulness.faithfulness_score": faithfulness.get("faithfulness_score", 0.0),
                "faithfulness.claim_count": faithfulness.get("claim_count", 0),
                "faithfulness.supported_claim_count": faithfulness.get("supported_claim_count", 0),
                "faithfulness.unsupported_claim_count": faithfulness.get("unsupported_claim_count", 0),
                "faithfulness.uncertain_claim_count": faithfulness.get("uncertain_claim_count", 0),
                "faithfulness.unsupported_claim_rate": faithfulness.get("unsupported_claim_rate", 0.0),
            }
        )
    if visual.get("status") == "ok":
        row.update(
            {
                "visual_readability.score": visual.get("score", 0.0),
                "visual_readability.rule_score": visual.get("rule_score", (visual.get("rule", {}) or {}).get("score", 0.0)),
                "visual_readability.vlm_score_combined_input": visual.get("vlm_score", (visual.get("vlm", {}) or {}).get("score", 0.0)),
                "visual_readability.text_overflow_count": visual.get("text_overflow_count", 0),
                "visual_readability.blank_space_count": visual.get("blank_space_count", 0),
                "visual_readability.missing_image_count": visual.get("missing_image_count", 0),
            }
        )
    if vlm_visual.get("status") == "ok":
        row.update(
            {
                "visual_readability.vlm_score": vlm_visual.get("score", 0.0),
                "visual_readability.vlm_aesthetic_average": vlm_visual.get("aesthetic_average", 0.0),
                "visual_readability.vlm_information_average": vlm_visual.get("information_average", 0.0),
            }
        )
    for criterion in VLM_VISUAL_CRITERIA:
        if vlm_visual.get("status") == "ok":
            row[f"visual_readability.vlm_{criterion}"] = (vlm_criteria.get(criterion, {}) or {}).get("score", 0.0)
    return row


def extract_json_object(text: str) -> dict:
    cleaned = text.strip()
    fenced = re.search(r"```(?:json)?\s*(\{.*?\})\s*```", cleaned, flags=re.S)
    if fenced:
        cleaned = fenced.group(1)
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        start = cleaned.find("{")
        end = cleaned.rfind("}")
        if start == -1 or end == -1 or end <= start:
            raise ValueError("LLM response did not contain a JSON object.")
        return json.loads(cleaned[start : end + 1])


def extract_qa_answers_from_text(text: str) -> dict[str, str]:
    answers: dict[str, str] = {}
    patterns = [
        r'"id"\s*:\s*"(?P<id>[^"]+)"[\s\S]{0,160}?"answer"\s*:\s*"?(?P<answer>[A-D])\b',
        r"(?P<id>(?:detail|understanding)_\d+)[^\n\rA-D]{0,80}\b(?P<answer>[A-D])\b",
    ]
    for pattern in patterns:
        for match in re.finditer(pattern, text, flags=re.I):
            answers[match.group("id")] = match.group("answer").upper()
    return answers


def normalize_claim_judgments(raw_claims: list[Any]) -> list[dict[str, str]]:
    claims = []
    valid_labels = {"supported", "unsupported", "uncertain"}
    for raw in raw_claims:
        if not isinstance(raw, dict):
            continue
        claim = normalize_space(str(raw.get("claim", "")))
        if not claim:
            continue
        label = str(raw.get("label", "")).strip().lower()
        if label not in valid_labels:
            label = "uncertain"
        claims.append(
            {
                "claim": claim,
                "label": label,
                "evidence": normalize_space(str(raw.get("evidence", ""))),
            }
        )
    return claims


def extract_claims_from_text(text: str) -> list[dict[str, str]]:
    claims = []
    json_like_pattern = (
        r'"claim"\s*:\s*"(?P<claim>(?:[^"\\]|\\.)*)"'
        r'[\s\S]{0,240}?'
        r'"label"\s*:\s*"?(?P<label>supported|unsupported|uncertain)\b'
    )
    for match in re.finditer(json_like_pattern, text, flags=re.I):
        claims.append(
            {
                "claim": match.group("claim").replace('\\"', '"').strip(),
                "label": match.group("label").lower(),
                "evidence": "Recovered from non-strict LLM output.",
            }
        )
    if claims:
        return claims
    pattern = r"(?P<label>supported|unsupported|uncertain)\s*[:\-]\s*(?P<claim>[^\n\r]+)"
    for match in re.finditer(pattern, text, flags=re.I):
        claims.append(
            {
                "claim": match.group("claim").strip(" -"),
                "label": match.group("label").lower(),
                "evidence": "Recovered from non-strict LLM output.",
            }
        )
    return claims


def answer_letter(value: str) -> str:
    match = re.search(r"\b([A-D])\b", value.upper())
    if match:
        return match.group(1)
    stripped = value.strip().upper()
    return stripped[0] if stripped[:1] in {"A", "B", "C", "D"} else ""


def read_json(path: Path) -> dict:
    if not path.is_file():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def normalize_space(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def count_words(value: str) -> int:
    return len(re.findall(r"[A-Za-z0-9]+(?:[-'][A-Za-z0-9]+)?", value))


def clamp(value: float, lower: float, upper: float) -> float:
    return max(lower, min(upper, value))


if __name__ == "__main__":
    main()
