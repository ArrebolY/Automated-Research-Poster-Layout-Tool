from __future__ import annotations

import math
import re
from pathlib import Path

from ..config import ExportConfig, LayoutConfig
from ..models import Panel, PosterLayout


class PptxExporter:
    def __init__(self, export_config: ExportConfig, layout_config: LayoutConfig) -> None:
        self.export_config = export_config
        self.layout_config = layout_config

    def export(self, layout: PosterLayout, output_path: Path) -> None:
        try:
            from PIL import Image
            from pptx import Presentation
            from pptx.dml.color import RGBColor
            from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
            from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
            from pptx.util import Inches, Pt
        except ModuleNotFoundError as exc:
            raise RuntimeError(
                "PPTX export requires the 'python-pptx' package. Install dependencies with `pip install -e .` first."
            ) from exc

        output_path.parent.mkdir(parents=True, exist_ok=True)

        prs = Presentation()
        prs.slide_width = Inches(self.layout_config.page_width)
        prs.slide_height = Inches(self.layout_config.page_height)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)

        self._add_title(slide, layout.title, Inches, Pt, PP_ALIGN)
        for panel in layout.panels:
            self._add_panel(
                slide,
                panel,
                Inches,
                Pt,
                MSO_AUTO_SHAPE_TYPE,
                RGBColor,
                Image,
                PP_ALIGN,
                MSO_AUTO_SIZE,
            )

        if len(prs.slides) > 1:
            prs.slides._sldIdLst.remove(prs.slides._sldIdLst[0])

        prs.save(output_path)

    def _add_title(self, slide, title: str, Inches, Pt, PP_ALIGN) -> None:
        left = Inches(0.4)
        top = Inches(0.15)
        width = Inches(self.layout_config.page_width - 0.8)
        height = Inches(0.5)
        box = slide.shapes.add_textbox(left, top, width, height)
        frame = box.text_frame
        paragraph = frame.paragraphs[0]
        paragraph.text = title
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.font.size = Pt(min(self.export_config.title_font_size_pt, 22))
        paragraph.font.bold = True

    def _add_panel(self, slide, panel: Panel, Inches, Pt, MSO_AUTO_SHAPE_TYPE, RGBColor, Image, PP_ALIGN, MSO_AUTO_SIZE) -> None:
        margin = self.layout_config.margin
        page_w = self.layout_config.page_width
        page_h = self.layout_config.page_height

        left = Inches((panel.x + margin) * page_w)
        top = Inches((panel.y + margin) * page_h)
        width = Inches(max(0.6, (panel.w - 2 * margin) * page_w))
        height = Inches(max(0.8, (panel.h - 2 * margin) * page_h))

        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(252, 253, 255)
        shape.line.color.rgb = RGBColor(176, 186, 198)
        shape.line.width = Pt(0.8)
        accent = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE,
            left,
            top,
            width,
            Inches(0.04),
        )
        accent.fill.solid()
        accent.fill.fore_color.rgb = RGBColor(30, 78, 121)
        accent.line.fill.background()

        primary_visual = self._primary_visual(panel)
        has_image = primary_visual is not None
        body_font_size = self._body_font_size(panel)
        if has_image:
            self._add_panel_with_visual(
                slide,
                panel,
                primary_visual,
                left,
                top,
                width,
                height,
                Inches,
                Pt,
                Image,
                PP_ALIGN,
                MSO_AUTO_SIZE,
                body_font_size,
            )
            return

        text_box = slide.shapes.add_textbox(
            left + Inches(0.12),
            top + Inches(0.1),
            width - Inches(0.24),
            height - Inches(0.18),
        )
        frame = text_box.text_frame
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        frame.margin_left = 0
        frame.margin_right = 0
        frame.margin_top = 0
        frame.margin_bottom = 0

        heading = frame.paragraphs[0]
        heading.text = panel.title
        heading.alignment = PP_ALIGN.LEFT
        heading.font.bold = True
        heading.font.color.rgb = RGBColor(30, 78, 121)
        heading.font.size = Pt(body_font_size + 2)

        fitted_body = self._fit_body_to_box(
            panel.title,
            panel.body,
            width - Inches(0.24),
            height - Inches(0.18),
            body_font_size,
            Inches,
        )
        self._add_body(frame, fitted_body, Pt, body_font_size, PP_ALIGN)
        self._add_figure_captions(frame, panel, Pt, PP_ALIGN)
        self._add_table_captions(frame, panel, Pt, PP_ALIGN)

    def _add_panel_with_visual(
        self,
        slide,
        panel: Panel,
        visual,
        left,
        top,
        width,
        height,
        Inches,
        Pt,
        Image,
        PP_ALIGN,
        MSO_AUTO_SIZE,
        body_font_size: int,
    ) -> None:
        mode = self._visual_layout_mode(panel, visual, Image)
        pad_x = Inches(0.12)
        pad_top = Inches(0.1)
        gap = Inches(0.08)
        if mode == "side_by_side":
            text_width = width * self._side_text_ratio(panel)
            visual_left = left + pad_x + text_width + gap
            visual_width = width - text_width - gap - Inches(0.2)
            text_box = slide.shapes.add_textbox(
                left + pad_x,
                top + pad_top,
                text_width,
                height - Inches(0.22),
            )
            self._fill_text_box(text_box, panel, Pt, body_font_size, PP_ALIGN, MSO_AUTO_SIZE, Inches)
            self._add_panel_image(
                slide,
                visual,
                visual_left,
                top + pad_top,
                visual_width,
                height - Inches(0.22),
                Inches,
                Pt,
                Image,
                PP_ALIGN,
            )
            return

        text_height = self._safe_stacked_text_height(panel, width, height, visual, body_font_size, Inches)
        stacked_gap = Inches(0.12)
        bottom_pad = Inches(0.06)
        text_box = slide.shapes.add_textbox(
            left + pad_x,
            top + pad_top,
            width - Inches(0.24),
            text_height,
        )
        self._fill_text_box(text_box, panel, Pt, body_font_size, PP_ALIGN, MSO_AUTO_SIZE, Inches)
        self._add_panel_image(
            slide,
            visual,
            left,
            top + pad_top + text_height + stacked_gap,
            width,
            height - pad_top - text_height - stacked_gap - bottom_pad,
            Inches,
            Pt,
            Image,
            PP_ALIGN,
        )

    def _fill_text_box(self, text_box, panel: Panel, Pt, body_font_size: int, PP_ALIGN, MSO_AUTO_SIZE, Inches) -> None:
        from pptx.dml.color import RGBColor

        frame = text_box.text_frame
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        frame.margin_left = 0
        frame.margin_right = 0
        frame.margin_top = 0
        frame.margin_bottom = 0

        heading = frame.paragraphs[0]
        heading.text = panel.title
        heading.alignment = PP_ALIGN.LEFT
        heading.font.bold = True
        heading.font.color.rgb = RGBColor(30, 78, 121)
        heading.font.size = Pt(body_font_size + 2)
        fitted_body = self._fit_body_to_box(
            panel.title,
            panel.body,
            text_box.width,
            text_box.height,
            body_font_size,
            Inches,
        )
        self._add_body(frame, fitted_body, Pt, body_font_size, PP_ALIGN)

    def _add_body(self, frame, body: str, Pt, font_size: int, PP_ALIGN) -> None:
        lines = [line.strip() for line in body.splitlines() if line.strip()]
        if lines and all(line.startswith("- ") for line in lines):
            for line in lines:
                paragraph = frame.add_paragraph()
                paragraph.text = line[2:].strip()
                paragraph.alignment = PP_ALIGN.LEFT
                paragraph.level = 0
                paragraph.font.size = Pt(font_size)
                paragraph._p.get_or_add_pPr().set("marL", "228600")
                paragraph._p.get_or_add_pPr().set("indent", "-114300")
        else:
            content = frame.add_paragraph()
            content.text = body
            content.alignment = PP_ALIGN.LEFT
            content.font.size = Pt(font_size)

    def _add_figure_captions(self, frame, panel: Panel, Pt, PP_ALIGN, exclude_id: str | None = None) -> None:
        for figure in panel.figures[:2]:
            if figure.figure_id == exclude_id:
                continue
            paragraph = frame.add_paragraph()
            paragraph.text = self._format_caption(figure.figure_id, figure.caption)
            paragraph.alignment = PP_ALIGN.LEFT
            paragraph.font.italic = True
            paragraph.font.size = Pt(self.export_config.caption_font_size_pt)

    def _add_table_captions(self, frame, panel: Panel, Pt, PP_ALIGN, exclude_id: str | None = None) -> None:
        for table in panel.tables[:2]:
            if table.table_id == exclude_id:
                continue
            paragraph = frame.add_paragraph()
            cell_note = f" ({len(table.cells)} rows)" if table.cells else ""
            paragraph.text = f"{self._format_caption(table.table_id, table.caption)}{cell_note}"
            paragraph.alignment = PP_ALIGN.LEFT
            paragraph.font.italic = True
            paragraph.font.size = Pt(7)

    def _add_panel_image(self, slide, visual, left, top, width, height, Inches, Pt, Image, PP_ALIGN) -> None:
        _kind, asset_id, caption, image_path, _aspect = visual
        image_left = left + Inches(0.2)
        image_top = top
        max_width = width - Inches(0.18)
        caption_height = self._caption_height(width, Inches)
        max_height = max(Inches(0.45), height - caption_height - Inches(0.1))
        image_width, image_height = self._fit_image(image_path, max_width, max_height, Image)
        image_left = left + (width - image_width) / 2
        slide.shapes.add_picture(str(image_path), image_left, image_top, width=image_width, height=image_height)

        caption_top = min(top + image_height + Inches(0.06), top + height - caption_height)
        caption_box = slide.shapes.add_textbox(left + Inches(0.09), caption_top, width - Inches(0.18), caption_height)
        caption_frame = caption_box.text_frame
        caption_frame.word_wrap = True
        caption_frame.margin_left = 0
        caption_frame.margin_right = 0
        caption_frame.margin_top = 0
        caption_frame.margin_bottom = 0
        paragraph = caption_frame.paragraphs[0]
        paragraph.text = self._format_caption(asset_id, caption)
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.font.italic = True
        paragraph.font.size = Pt(7)

    def _primary_visual(self, panel: Panel):
        candidates = []
        for figure in panel.figures:
            if figure.path is not None:
                aspect = self._image_aspect(figure.path)
                candidates.append(
                    (figure.match_score, "figure", figure.figure_id, figure.caption, figure.path, aspect)
                )
        for table in panel.tables:
            if table.path is not None:
                if table.cells and len(table.cells) > 20:
                    continue
                row_penalty = max(0, len(table.cells) - 12) * 4.0
                aspect = self._image_aspect(table.path)
                candidates.append(
                    (table.match_score - row_penalty, "table", table.table_id, table.caption, table.path, aspect)
                )
        if candidates:
            _score, kind, asset_id, caption, path, aspect = max(candidates, key=lambda item: item[0])
            return kind, asset_id, caption, path, aspect
        return None

    def _visual_layout_mode(self, panel: Panel, visual, Image) -> str:
        kind, _asset_id, _caption, image_path, aspect = visual
        panel_aspect = (panel.w * self.layout_config.page_width) / max(0.01, panel.h * self.layout_config.page_height)
        visual_aspect = aspect or self._image_aspect(image_path) or 1.0
        word_count = len(panel.body.split())
        if kind == "table":
            return "stacked"
        title_lines = max(1, math.ceil(len(panel.title) / 34))
        estimated_text_lines = title_lines + self._estimate_body_lines(panel.body)
        panel_height_in = panel.h * self.layout_config.page_height
        if (
            panel_aspect >= 2.15
            and panel_height_in >= 1.75
            and visual_aspect <= 3.2
            and word_count <= 75
            and estimated_text_lines <= 7
        ):
            return "side_by_side"
        return "stacked"

    def _side_text_ratio(self, panel: Panel) -> float:
        word_count = len(panel.body.split())
        if word_count <= 45:
            return 0.42
        if word_count <= 75:
            return 0.48
        return 0.54

    def _text_height(self, panel: Panel, height, has_image: bool, visual=None):
        if not has_image:
            return height * 0.9
        word_count = len(panel.body.split())
        kind = visual[0] if visual else "figure"
        if kind == "table":
            ratio = 0.28 if word_count <= 55 else 0.34
        else:
            if word_count <= 45:
                ratio = 0.25
            elif word_count <= 75:
                ratio = 0.32
            else:
                ratio = 0.4
        return max(height * 0.18, min(height * ratio, height * 0.46))

    def _safe_stacked_text_height(self, panel: Panel, width, height, visual, body_font_size: int, Inches):
        estimated = self._estimate_text_height(panel, body_font_size, Inches, width)
        base = self._text_height(panel, height, True, visual)
        minimum_image_height = Inches(0.85)
        reserved_spacing = Inches(0.1) + Inches(0.16) + Inches(0.08)
        max_text_height = max(Inches(0.46), height - reserved_spacing - minimum_image_height)
        return min(max(base, estimated), max_text_height)

    def _estimate_text_height(self, panel: Panel, body_font_size: int, Inches, width=None):
        title_lines = max(1, math.ceil(len(panel.title) / 34))
        body_lines = self._estimate_body_lines(panel.body, self._words_per_line(width, Inches) if width else 11)
        line_height = (body_font_size + 3) / 72
        return Inches(max(0.42, (title_lines + body_lines) * line_height + 0.08))

    def _estimate_body_lines(self, body: str, words_per_line: int = 11) -> int:
        lines = [line.strip() for line in body.splitlines() if line.strip()]
        if lines:
            return sum(max(1, math.ceil(len(line.split()) / words_per_line)) for line in lines[:6])
        return max(1, math.ceil(len(body.split()) / words_per_line))

    def _words_per_line(self, width, Inches) -> int:
        width_in = max(1.0, width / Inches(1) - 0.35)
        return max(10, min(24, int(width_in * 3.6)))

    def _fit_body_to_box(self, title: str, body: str, width, height, font_size: int, Inches) -> str:
        width_in = max(0.8, width / Inches(1))
        height_in = max(0.3, height / Inches(1))
        title_chars_per_line = max(16, int(width_in * 10))
        title_lines = max(1, math.ceil(len(title) / title_chars_per_line))
        heading_points = title_lines * (font_size + 2) * 1.22
        body_line_points = (font_size + 2) * 1.18
        available_points = max(0, height_in * 72 - heading_points - 6)
        available_lines = max(1, int(available_points / body_line_points))
        max_words = max(6, int(available_lines * self._words_per_line(width, Inches) * 0.78))
        if self._word_count(body) <= max_words:
            return body
        return self._trim_body_to_words(body, max_words)

    def _trim_body_to_words(self, body: str, max_words: int) -> str:
        lines = [line.strip() for line in body.splitlines() if line.strip()]
        if lines and all(line.startswith("- ") for line in lines):
            result: list[str] = []
            used = 0
            for line in lines:
                words = line[2:].split()
                remaining = max_words - used
                if remaining <= 3:
                    break
                take = min(len(words), remaining, 22)
                result.append("- " + self._complete_phrase(words, take))
                used += take
            return "\n".join(result or lines[:1])
        words = body.split()
        return self._complete_phrase(words, max_words)

    def _word_count(self, text: str) -> int:
        return len([word for word in text.split() if word.strip()])

    def _caption_height(self, width, Inches):
        width_in = width / Inches(1)
        return Inches(0.22 if width_in < 3.0 else 0.16)

    def _format_caption(self, asset_id: str, caption: str) -> str:
        if re.match(rf"^\s*{re.escape(asset_id)}\s*:", caption.strip(), flags=re.IGNORECASE):
            return self._trim_caption(caption.strip())
        clean_caption = self._strip_caption_label(caption)
        return f"{asset_id}: {self._trim_caption(clean_caption)}"

    def _strip_caption_label(self, caption: str) -> str:
        import re

        return re.sub(
            r"^\s*(?:fig(?:ure)?\.?|table)\s*\d+[a-z]?\s*[.:-]\s*",
            "",
            caption,
            flags=re.IGNORECASE,
        ).strip()

    def _trim_caption(self, caption: str) -> str:
        words = caption.split()
        if len(words) <= 18:
            return caption
        import re

        first_sentence = re.split(r"(?<=[.!?])\s+", caption.strip())[0].strip()
        if 4 <= len(first_sentence.split()) <= 18:
            return first_sentence.rstrip(" .,;:")
        return self._complete_caption_phrase(words, 18)

    def _complete_caption_phrase(self, words: list[str], max_words: int) -> str:
        stop_endings = {"and", "or", "with", "by", "to", "from", "for", "of", "in", "as", "than", "while", "under", "over"}
        chosen = words[:max_words]
        while len(chosen) > 4 and chosen[-1].strip(".,;:").lower() in stop_endings:
            chosen = chosen[:-1]
        return " ".join(chosen).rstrip(" .,;:")

    def _complete_phrase(self, words: list[str], max_words: int) -> str:
        if not words:
            return ""
        text = " ".join(words[:max_words]).strip()
        text = text.removesuffix("...").rstrip(" .,;:")
        if not text.endswith((".", "!", "?")):
            text += "."
        return text

    def _fit_image(self, image_path: Path, max_width, max_height, Image):
        with Image.open(image_path) as image:
            ratio = image.width / image.height if image.height else 1

        width = max_width
        height = int(width / ratio)
        if height > max_height:
            height = max_height
            width = int(height * ratio)
        return width, height

    def _image_aspect(self, image_path: Path) -> float | None:
        try:
            from PIL import Image

            with Image.open(image_path) as image:
                return image.width / image.height if image.height else None
        except (OSError, ModuleNotFoundError):
            return None

    def _body_font_size(self, panel: Panel) -> int:
        area = panel.w * panel.h
        if area < 0.18:
            return 8
        if panel.figures or panel.tables:
            return 9
        return 11
