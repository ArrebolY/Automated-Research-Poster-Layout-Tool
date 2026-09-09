import json
from pathlib import Path

from pptx import Presentation

from postergen.final_eval import (
    MethodSpec,
    VLMConfig,
    answer_letter,
    evaluate_final_poster,
    evaluate_paperquiz_vlm,
    evaluate_visual_readability,
    extract_claims_from_text,
    extract_html_text,
    extract_pptx_text,
    extract_qa_answers_from_text,
    find_method_artifact,
    flatten_final_result,
    normalize_vlm_judge_result,
    normalize_questions,
    parse_methods,
)


class FakeGenerator:
    def generate(self, prompt: str) -> str:
        if "PaperQuiz" in prompt:
            return json.dumps(
                {
                    "answers": [
                        {"id": "detail_1", "answer": "A", "confidence": 1.0, "evidence": "title"},
                        {"id": "understanding_1", "answer": "B", "confidence": 1.0, "evidence": "method"},
                    ]
                }
            )
        return json.dumps(
            {
                "claims": [
                    {"claim": "Paper A presents a method.", "label": "supported", "evidence": "method"},
                    {"claim": "Paper A reports a result.", "label": "supported", "evidence": "result"},
                    {"claim": "Paper A uses a different dataset.", "label": "unsupported", "evidence": "not in paper"},
                ]
            }
        )


def test_extract_poster_text_from_pptx_and_html(tmp_path):
    pptx_path = tmp_path / "poster.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, 1000000, 1000000)
    box.text = "Poster title\nMain result"
    prs.save(pptx_path)

    html_path = tmp_path / "poster.html"
    html_path.write_text("<html><style>.a{}</style><body><h1>Poster</h1><p>Text &amp; data</p></body></html>")

    assert "Main result" in extract_pptx_text(pptx_path)
    assert extract_html_text(html_path) == "Poster Text & data"


def test_normalize_questions_supports_paperquiz_dict_shape():
    block = {
        "questions": {
            "Question 1": {"question": "Title?", "options": ["A. Good", "B. Bad"]},
        },
        "answers": {"Question 1": "A. Good"},
        "aspects": {"Question 1": "title"},
    }

    questions = normalize_questions(block, "detail", 10)

    assert questions[0].id == "detail_1"
    assert questions[0].options == ["A. Good", "B. Bad"]
    assert answer_letter(questions[0].answer) == "A"


def test_recover_answers_and_claims_from_non_strict_llm_output():
    bad_qa = '{"answers":[{"id":"detail_1","answer":"A","evidence":"poster says "quoted" text"}]}'
    bad_claims = """
    {"claims":[
      {"claim":"The method is evaluated on a benchmark with \"quoted\" text","label":"supported"},
      {"claim":"The poster claims a 20% gain that is not in the paper","label":"unsupported"}
    ]}
    """

    assert extract_qa_answers_from_text(bad_qa)["detail_1"] == "A"
    claims = extract_claims_from_text(bad_claims)
    assert claims[0]["label"] == "supported"
    assert claims[0]["claim"].startswith("The method is evaluated")
    assert claims[1]["label"] == "unsupported"


def test_find_method_artifact_handles_batch_mode_stems(tmp_path):
    paper = tmp_path / "Sample Paper" / "paper.pdf"
    paper.parent.mkdir()
    paper.write_bytes(b"%PDF")
    root = tmp_path
    method_root = tmp_path / "outputs" / "batch_posters" / "template_based"
    method_root.mkdir(parents=True)
    expected = method_root / "Sample_Paper_template_based.pptx"
    expected.write_bytes(b"pptx")

    found = find_method_artifact(MethodSpec("template_based", method_root), paper, root)

    assert found == expected


def test_default_proposed_method_uses_proposed_batch_posters_root():
    assert parse_methods([])[0] == MethodSpec("proposed", Path("outputs/batch_posters/proposed"))


def test_evaluate_final_poster_combines_qa_faithfulness_and_visual(tmp_path):
    paper = tmp_path / "Paper A" / "paper.pdf"
    paper.parent.mkdir()
    document = Presentation()
    slide = document.slides.add_slide(document.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, 1000000, 1000000).text = "Paper A method result conclusion"
    artifact = tmp_path / "poster.pptx"
    document.save(artifact)
    qa = {
        "detail": {
            "questions": {"Question 1": {"question": "Title?", "options": ["A. Paper A", "B. Other"]}},
            "answers": {"Question 1": "A"},
            "aspects": {"Question 1": "title"},
        },
        "understanding": {
            "questions": {"Question 1": {"question": "Method?", "options": ["A. Wrong", "B. Method"]}},
            "answers": {"Question 1": "B"},
            "aspects": {"Question 1": "method"},
        },
    }
    (paper.parent / "o3_qa.json").write_text(json.dumps(qa), encoding="utf-8")
    paper.write_bytes(b"%PDF-1.4\n")
    report_dir = tmp_path / "assets" / "Paper_A"
    report_dir.mkdir(parents=True)
    (report_dir / "paper.md").write_text("Paper A method result conclusion", encoding="utf-8")
    (report_dir / "panel_quality.json").write_text(
        json.dumps({"summary": {"issue_counts": {}}, "panels": [{"text_occupancy": 0.8}]}),
        encoding="utf-8",
    )
    (report_dir / "layout_quality.json").write_text(json.dumps({"figures": [{"id": "Figure 1"}]}), encoding="utf-8")

    result = evaluate_final_poster(
        paper_path=paper,
        method="proposed",
        artifact=artifact,
        report_dir=report_dir,
        generator=FakeGenerator(),
        max_questions_per_type=10,
        paperquiz_evaluator="text",
    )

    assert result["paperquiz"]["average_accuracy"] == 1.0
    assert result["faithfulness"]["faithfulness_score"] == 66.667
    assert result["faithfulness"]["unsupported_claim_rate"] == 0.333
    assert result["visual_readability"]["score"] == 5.0
    assert result["overall"]["overall_score"] > 80


def test_evaluate_paperquiz_vlm_answers_from_rendered_image(monkeypatch, tmp_path):
    paper = tmp_path / "Paper A"
    paper.mkdir()
    qa = {
        "detail": {
            "questions": {"Question 1": {"question": "Title?", "options": ["A. Paper A", "B. Other"]}},
            "answers": {"Question 1": "A"},
            "aspects": {"Question 1": "title"},
        },
        "understanding": {
            "questions": {"Question 1": {"question": "Method?", "options": ["A. Wrong", "B. Method"]}},
            "answers": {"Question 1": "B"},
            "aspects": {"Question 1": "method"},
        },
    }
    qa_path = paper / "o3_qa.json"
    qa_path.write_text(json.dumps(qa), encoding="utf-8")
    artifact = tmp_path / "poster.pptx"
    artifact.write_bytes(b"pptx")

    def fake_render(source, output):
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_bytes(b"png")
        return {"status": "ok", "backend": "fake", "image_path": str(output)}

    def fake_answer(image_path, questions_payload, config):
        answers = []
        for question in questions_payload:
            answers.append({"id": question["id"], "answer": "A" if question["id"].startswith("detail") else "B"})
        return {"answers": answers}

    monkeypatch.setattr("postergen.final_eval.render_artifact_to_png", fake_render)
    monkeypatch.setattr("postergen.final_eval.openai_vlm_answer_questions", fake_answer)

    result = evaluate_paperquiz_vlm(
        qa_path,
        artifact,
        method="proposed",
        paper_id="Paper A",
        max_questions_per_type=10,
        vlm_config=VLMConfig("openai", "gpt-4o-mini", "OPENAI_API_KEY"),
        poster_images_root=tmp_path / "images",
    )

    assert result["status"] == "ok"
    assert result["average_accuracy"] == 1.0
    assert result["artifact_image"].endswith(".png")


def test_normalize_vlm_judge_result_averages_criteria():
    data = {
        "criteria": {
            "aesthetic_element": {"score": 4, "reason": "balanced"},
            "aesthetic_engagement": {"score": 3, "reason": "plain"},
            "aesthetic_layout": {"score": 5, "reason": "clear"},
            "information_low_level": {"score": 4, "reason": "readable"},
            "information_logic": {"score": 4, "reason": "ordered"},
            "information_content": {"score": 4, "reason": "relevant"},
        },
        "overall": {"score": 4, "reason": "good"},
    }

    result = normalize_vlm_judge_result(data)

    assert result["score"] == 4.0
    assert result["aesthetic_average"] == 4.0
    assert result["information_average"] == 4.0
    assert result["criteria"]["aesthetic_layout"]["score"] == 5


def test_visual_readability_vlm_uses_rendered_image_and_flattens(monkeypatch, tmp_path):
    artifact = tmp_path / "poster.html"
    artifact.write_text("<h1>Poster</h1>", encoding="utf-8")

    def fake_render(source, output):
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_bytes(b"png")
        return {"status": "ok", "backend": "fake", "image_path": str(output)}

    def fake_judge(image_path, config):
        return {
            "score": 4.2,
            "criteria": {
                "aesthetic_element": {"score": 4, "reason": ""},
                "aesthetic_engagement": {"score": 4, "reason": ""},
                "aesthetic_layout": {"score": 4, "reason": ""},
                "information_low_level": {"score": 5, "reason": ""},
                "information_logic": {"score": 4, "reason": ""},
                "information_content": {"score": 4, "reason": ""},
            },
            "overall": {"score": 4, "reason": ""},
        }

    monkeypatch.setattr("postergen.final_eval.render_artifact_to_png", fake_render)
    monkeypatch.setattr("postergen.final_eval.openai_vlm_judge", fake_judge)

    visual = evaluate_visual_readability(
        artifact,
        tmp_path,
        "Poster text",
        method="direct_llm",
        paper_id="Paper",
        visual_evaluator="both",
        vlm_config=VLMConfig("openai", "gpt-4o-mini", "OPENAI_API_KEY"),
        poster_images_root=tmp_path / "images",
    )
    row = flatten_final_result(
        {
            "paper_id": "Paper",
            "method": "direct_llm",
            "artifact": str(artifact),
            "artifact_exists": True,
            "paperquiz": {},
            "faithfulness": {},
            "visual_readability": visual,
            "overall": {},
        }
    )

    assert visual["score"] == 3.81
    assert visual["rule_score"] == 2.9
    assert visual["vlm_score"] == 4.2
    assert visual["score_source"] == "vlm_rule_combined"
    assert row["visual_readability.vlm_information_low_level"] == 5
