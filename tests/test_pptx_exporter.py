from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN

from postergen.config import ExportConfig, LayoutConfig
from postergen.models import Figure, Panel, PosterLayout, Table
from postergen.stages.pptx_exporter import PptxExporter


def test_pptx_exporter_renders_table_image_and_removes_duplicate_caption_label(tmp_path):
    table_path = tmp_path / "table.png"
    Image.new("RGB", (500, 240), "white").save(table_path)
    table = Table(
        table_id="Table 1",
        caption="Table 1. Evaluation results across models.",
        cells=[["Model", "Score"], ["Ours", "90"]],
        path=table_path,
    )
    panel = Panel(
        title="Results",
        body="- The proposed system achieves the strongest score.",
        figures=[],
        tables=[table],
        importance=1.0,
        x=0.0,
        y=0.12,
        w=1.0,
        h=0.88,
    )
    output = tmp_path / "poster.pptx"

    PptxExporter(ExportConfig(), LayoutConfig()).export(PosterLayout("Demo", [panel]), output)

    presentation = Presentation(output)
    shapes = presentation.slides[0].shapes
    text = "\n".join(shape.text for shape in shapes if hasattr(shape, "text"))
    assert sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in shapes) == 1
    assert "Table 1: Evaluation results" in text
    assert "Table 1: Table 1" not in text


def test_pptx_exporter_prefers_figure_and_left_aligns_panel_text(tmp_path):
    figure_path = tmp_path / "figure.png"
    table_path = tmp_path / "table.png"
    Image.new("RGB", (600, 260), "white").save(figure_path)
    Image.new("RGB", (500, 900), "white").save(table_path)
    panel = Panel(
        title="Results",
        body="- The method improves quality.\n- Human evaluation confirms the result.",
        figures=[
            Figure(
                figure_id="Figure 3",
                caption="Overview of the evaluation results.",
                path=figure_path,
                match_score=8.0,
            )
        ],
        tables=[
            Table(
                table_id="Table 1",
                caption="Detailed results for all models.",
                cells=[["Model", "Score"]] * 39,
                path=table_path,
                match_score=100.0,
            )
        ],
        importance=1.0,
        x=0.0,
        y=0.12,
        w=1.0,
        h=0.88,
    )
    output = tmp_path / "poster.pptx"

    PptxExporter(ExportConfig(), LayoutConfig()).export(PosterLayout("Demo", [panel]), output)

    presentation = Presentation(output)
    shapes = presentation.slides[0].shapes
    pictures = [shape for shape in shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    panel_text = next(shape for shape in shapes if hasattr(shape, "text") and "Results" in shape.text)
    assert len(pictures) == 1
    assert round(pictures[0].width / pictures[0].height, 1) == round(600 / 260, 1)
    assert all(
        paragraph.alignment == PP_ALIGN.LEFT
        for paragraph in panel_text.text_frame.paragraphs
        if paragraph.text
    )


def test_pptx_exporter_does_not_render_dense_table_as_full_image(tmp_path):
    table_path = tmp_path / "table.png"
    Image.new("RGB", (500, 900), "white").save(table_path)
    table = Table(
        table_id="Table 1",
        caption="Detailed results for all evaluated models.",
        cells=[["Model", "Score"]] * 39,
        path=table_path,
        match_score=100.0,
    )
    panel = Panel(
        title="Results",
        body="- The proposed system achieves the strongest score.",
        figures=[],
        tables=[table],
        importance=1.0,
        x=0.0,
        y=0.12,
        w=1.0,
        h=0.88,
    )
    output = tmp_path / "poster.pptx"

    PptxExporter(ExportConfig(), LayoutConfig()).export(PosterLayout("Demo", [panel]), output)

    presentation = Presentation(output)
    shapes = presentation.slides[0].shapes
    text = "\n".join(shape.text for shape in shapes if hasattr(shape, "text"))
    assert not any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in shapes)
    assert "Table 1: Detailed results" in text
    assert "(39 rows)" in text


def test_pptx_exporter_uses_side_by_side_for_wide_panel_with_compact_figure(tmp_path):
    figure_path = tmp_path / "figure.png"
    Image.new("RGB", (360, 300), "white").save(figure_path)
    panel = Panel(
        title="Calibration",
        body="- The plot shows better calibrated predictions.\n- Lower uncertainty corresponds to higher accuracy.",
        figures=[
            Figure(
                figure_id="Figure 8",
                caption="Calibration plot for uncertainty estimates.",
                path=figure_path,
                match_score=10.0,
            )
        ],
        tables=[],
        importance=1.0,
        x=0.0,
        y=0.12,
        w=1.0,
        h=0.28,
    )
    output = tmp_path / "poster.pptx"

    PptxExporter(ExportConfig(), LayoutConfig()).export(PosterLayout("Demo", [panel]), output)

    presentation = Presentation(output)
    pictures = [shape for shape in presentation.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    text_boxes = [shape for shape in presentation.slides[0].shapes if hasattr(shape, "text") and "Calibration" in shape.text]
    assert len(pictures) == 1
    assert text_boxes
    assert pictures[0].left > text_boxes[0].left + text_boxes[0].width


def test_pptx_exporter_keeps_stacked_text_above_image(tmp_path):
    figure_path = tmp_path / "wide_figure.png"
    Image.new("RGB", (1100, 220), "white").save(figure_path)
    panel = Panel(
        title="Introduction",
        body=(
            "Uncertainty quantification is formulated as computing a rate-distortion function "
            "for a compressed training dataset representation.\n"
            "A meta-probabilistic perspective regularizes representations and enables distance awareness."
        ),
        figures=[
            Figure(
                figure_id="Figure 1",
                caption="Distance awareness for principled uncertainty quantification.",
                path=figure_path,
                match_score=10.0,
            )
        ],
        tables=[],
        importance=1.0,
        x=0.0,
        y=0.12,
        w=0.58,
        h=0.34,
    )
    output = tmp_path / "poster.pptx"

    PptxExporter(ExportConfig(), LayoutConfig()).export(PosterLayout("Demo", [panel]), output)

    presentation = Presentation(output)
    shapes = presentation.slides[0].shapes
    picture = next(shape for shape in shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
    text_box = next(shape for shape in shapes if hasattr(shape, "text") and "Introduction" in shape.text)
    assert picture.top > text_box.top + text_box.height


def test_pptx_exporter_wraps_visual_caption_inside_panel(tmp_path):
    figure_path = tmp_path / "figure.png"
    Image.new("RGB", (700, 360), "white").save(figure_path)
    panel = Panel(
        title="Introduction",
        body="- The visual summarizes the central idea.",
        figures=[
            Figure(
                figure_id="Figure 1",
                caption=(
                    "Distance awareness for principled uncertainty quantification. "
                    "A distance-aware representation separates correct and wrong predictions."
                ),
                path=figure_path,
                match_score=10.0,
            )
        ],
        tables=[],
        importance=1.0,
        x=0.0,
        y=0.12,
        w=0.58,
        h=0.34,
    )
    output = tmp_path / "poster.pptx"

    PptxExporter(ExportConfig(), LayoutConfig()).export(PosterLayout("Demo", [panel]), output)

    presentation = Presentation(output)
    caption_box = next(shape for shape in presentation.slides[0].shapes if hasattr(shape, "text") and "Figure 1" in shape.text)
    assert caption_box.text_frame.word_wrap
    assert caption_box.width < presentation.slide_width


def test_pptx_exporter_trims_text_to_available_box(tmp_path):
    panel = Panel(
        title="Conclusion",
        body=" ".join(["overflowing"] * 180),
        figures=[],
        tables=[],
        importance=1.0,
        x=0.0,
        y=0.12,
        w=0.45,
        h=0.18,
    )
    output = tmp_path / "poster.pptx"

    PptxExporter(ExportConfig(), LayoutConfig()).export(PosterLayout("Demo", [panel]), output)

    presentation = Presentation(output)
    text = "\n".join(shape.text for shape in presentation.slides[0].shapes if hasattr(shape, "text"))
    assert text.count("overflowing") < 180
    assert "..." not in text
    assert text.rstrip().endswith(".")
